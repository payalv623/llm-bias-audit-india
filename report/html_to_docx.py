"""
html_to_docx.py -- convert the generated report HTML (Assignment3_Report.html,
Assignment3_Short_Report.html) into Word .docx files with python-docx.

The HTML is produced by build_report.py / build_short_report.py, so only the
tags those scripts emit are handled: headings, paragraphs, lists, tables,
figures (img + figcaption), <pre> blocks and inline b/i/sub/sup/code/br.
Figure widths follow the CSS width (%) used in the HTML; the footer carries
the page number only.

Run from the project root:   .venv/bin/python report/html_to_docx.py
"""
import re
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
MARGIN_MM = 15
TEXT_WIDTH_MM = 210 - 2 * MARGIN_MM          # A4 width minus left/right margins
BLOCK_TAGS = {"p", "div", "ul", "ol", "table", "figure", "pre", "h1", "h2", "h3", "h4", "li"}


# ---------------------------------------------------------------------------
# small python-docx helpers
# ---------------------------------------------------------------------------
def add_page_number_footer(section):
    """Footer = centred PAGE field, nothing else."""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}),
                      ("w:instrText", None),
                      ("w:fldChar", {"w:fldCharType": "end"})):
        el = OxmlElement(tag)
        if attr:
            for k, v in attr.items():
                el.set(qn(k), v)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        run._r.append(el)
    run.font.size = Pt(9)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def clean(text):
    """Collapse HTML whitespace (newlines inside the source) to single spaces."""
    return re.sub(r"\s+", " ", text)


# ---------------------------------------------------------------------------
# inline content -> runs
# ---------------------------------------------------------------------------
def add_inline(par, node, fmt=None, size=None):
    """Append the inline content of `node` to paragraph `par`, keeping b/i/sub/sup/code."""
    fmt = dict(fmt or {})
    for ch in node.children:
        if isinstance(ch, NavigableString):
            t = clean(str(ch))
            if not t:
                continue
            r = par.add_run(t)
            r.bold = fmt.get("b")
            r.italic = fmt.get("i")
            r.font.subscript = fmt.get("sub")
            r.font.superscript = fmt.get("sup")
            if fmt.get("code"):
                r.font.name = "Consolas"
            if size:
                r.font.size = Pt(size)
            continue
        name = ch.name
        if name == "br":
            par.add_run().add_break()
            continue
        sub = dict(fmt)
        if name in ("b", "strong", "th"):
            sub["b"] = True
        elif name in ("i", "em"):
            sub["i"] = True
        elif name == "sub":
            sub["sub"] = True
        elif name == "sup":
            sub["sup"] = True
        elif name == "code":
            sub["code"] = True
        add_inline(par, ch, sub, size)


def strip_edges(par):
    """Trim leading/trailing spaces left by HTML indentation."""
    runs = par.runs
    if runs:
        runs[0].text = runs[0].text.lstrip()
        runs[-1].text = runs[-1].text.rstrip()


def has_block_children(el):
    return any(getattr(c, "name", None) in BLOCK_TAGS for c in el.children)


# ---------------------------------------------------------------------------
# block content
# ---------------------------------------------------------------------------
class Converter:
    def __init__(self, base_size):
        self.doc = Document()
        self.base = base_size
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(base_size)
        st.paragraph_format.space_after = Pt(3)
        for lvl, sz in ((1, 16), (2, 13), (3, 11.5), (4, 10.5)):
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.size = Pt(sz)
            h.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)
            h.font.name = "Calibri"
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Mm(MARGIN_MM))
        add_page_number_footer(sec)
        self.first_h2 = True

    # -- dispatch -----------------------------------------------------------
    def block(self, el, list_level=0):
        name = el.name
        cls = el.get("class") or []
        if name in ("h1", "h2", "h3", "h4"):
            if name == "h2" and "pb" in cls and not self.first_h2:
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if name == "h2":
                self.first_h2 = False
            lvl = int(name[1])
            h = self.doc.add_heading(level=lvl)
            add_inline(h, el)
            strip_edges(h)
            if lvl == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif name in ("ul", "ol"):
            for li in el.find_all("li", recursive=False):
                self.list_item(li, ordered=(name == "ol"), level=list_level)
        elif name == "table":
            self.table(el)
        elif name == "figure":
            self.figure(el)
        elif name == "pre":
            p = self.doc.add_paragraph()
            r = p.add_run(el.get_text())
            r.font.name = "Consolas"
            r.font.size = Pt(self.base - 1.5)
            p.paragraph_format.left_indent = Mm(4)
        elif name in ("div", "p", "span") and has_block_children(el):
            # container: walk children, gathering loose inline text into paragraphs
            self.container(el, list_level)
        elif name in ("p", "div", "span"):
            size = self.base - 2 if "plog" in cls or "small" in cls else None
            p = self.doc.add_paragraph()
            add_inline(p, el, size=size)
            strip_edges(p)
            if not p.text.strip():
                p._element.getparent().remove(p._element)
                return
            if "meta" in cls or "sub" in cls:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if "box" in cls:
                p.paragraph_format.left_indent = Mm(4)

    def container(self, el, list_level=0):
        buf = []

        def flush():
            if buf:
                tmp = BeautifulSoup("<p></p>", "html.parser").p
                for b in buf:
                    tmp.append(b)
                self.block(tmp)
                buf.clear()

        for ch in list(el.children):
            if getattr(ch, "name", None) in BLOCK_TAGS:
                flush()
                self.block(ch, list_level)
            else:
                buf.append(ch.__copy__() if hasattr(ch, "__copy__") else ch)
        flush()

    def list_item(self, li, ordered, level):
        style = ("List Number" if ordered else "List Bullet") + ("" if level == 0 else " 2")
        p = self.doc.add_paragraph(style=style)
        nested = []
        for ch in li.children:
            if getattr(ch, "name", None) in ("ul", "ol"):
                nested.append(ch)
            else:
                wrap = BeautifulSoup("<span></span>", "html.parser").span
                wrap.append(ch.__copy__() if hasattr(ch, "__copy__") else ch)
                add_inline(p, wrap)
        strip_edges(p)
        for n in nested:
            self.block(n, list_level=level + 1)

    def table(self, el):
        rows = el.find_all("tr")
        if not rows:
            return
        ncol = max(sum(int(c.get("colspan", 1)) for c in r.find_all(["td", "th"])) for r in rows)
        t = self.doc.add_table(rows=0, cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        size = self.base - 1.5 if ncol <= 6 else self.base - 2.5
        for r in rows:
            cells = r.find_all(["td", "th"])
            row = t.add_row().cells
            j = 0
            for c in cells:
                span = int(c.get("colspan", 1))
                cell = row[j]
                if span > 1:
                    cell = cell.merge(row[j + span - 1])
                par = cell.paragraphs[0]
                add_inline(par, c, {"b": c.name == "th"}, size=size)
                strip_edges(par)
                par.paragraph_format.space_after = Pt(0)
                if c.name == "th":
                    shade(cell, "E4E3DF")
                j += span
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, el):
        img = el.find("img")
        if img is not None:
            m = re.search(r"width:\s*([\d.]+)%", img.get("style", ""))
            frac = float(m.group(1)) / 100 if m else 1.0
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(HERE / img["src"]), width=Mm(TEXT_WIDTH_MM * frac))
        cap = el.find("figcaption")
        if cap is not None:
            p = self.doc.add_paragraph()
            add_inline(p, cap, size=self.base - 1.5)
            strip_edges(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def convert(self, html_path, out_path):
        soup = BeautifulSoup(Path(html_path).read_text(encoding="utf-8"), "html.parser")
        for ch in soup.body.children:
            if getattr(ch, "name", None):
                self.block(ch)
        self.doc.core_properties.author = "Payal Verma"
        self.doc.save(out_path)
        print("DOCX written:", out_path)


if __name__ == "__main__":
    Converter(10).convert(HERE / "Assignment3_Short_Report.html", HERE / "Assignment3_Short_Report.docx")
    Converter(10.5).convert(HERE / "Assignment3_Report.html", HERE / "Assignment3_Report.docx")
